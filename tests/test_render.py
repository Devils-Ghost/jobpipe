import copy

from docx import Document
from docx.oxml.ns import qn

from jobpipe import render
from jobpipe.render import Adjust, DocStats, drop_lowest_bullet, drop_lowest_project, fit_one_page, ladder
from jobpipe.render.template import master_context

FLOORS = {"paragraph_spacing_pt": 0, "line_spacing": 1.0, "margin_in": 0.5, "body_font_pt": 10.0}


def _ctx(project_bullets=(2, 2, 2)):
    seg = lambda t: [{"t": t, "b": False}]
    return {
        "skills": {"languages": seg("x")},
        "projects": [
            {"title": f"P{j}", "term": "2025", "before": 60, "ind_right": 0, "num_id": "5",
             "bullets": [{"segs": seg(f"p{j}-{i}"), "ind_right": 0} for i in range(n)]}
            for j, n in enumerate(project_bullets)
        ],
    }


def _texts(f):
    ps = Document(str(f)).element.body.iter(qn("w:p"))
    return [t for t in ("".join(x.text or "" for x in p.iter(qn("w:t"))) for p in ps) if t.strip()]


def _links(docx_path):
    doc = Document(str(docx_path))
    return {
        "".join(t.text or "" for t in h.iter(qn("w:t"))): doc.part.rels[h.get(qn("r:id"))].target_ref
        for h in doc.element.body.iter(qn("w:hyperlink"))
    }


# ---------------------------------------------------------------- ladder order


def test_ladder_full_order():
    stats = DocStats(spacing=True, line_spacing=1.15, margin_in=0.75, font_pt=11.0)
    steps = [d for d, _ in ladder(stats, FLOORS, 0.5)]
    assert steps == ["paragraph spacing -> 0", "line spacing -> 1.0", "margins -> 0.5 in",
                     "body font -> 10.5 pt", "body font -> 10.0 pt"]
    *_, (_, final) = ladder(stats, FLOORS, 0.5)
    assert final == Adjust(zero_spacing=True, line_spacing=1.0, margin_in=0.5, font_pt=10.0)


def test_ladder_skips_noops_for_dense_resume():
    stats = DocStats(spacing=True, line_spacing=1.0, margin_in=0.3, font_pt=10.0)
    assert [d for d, _ in ladder(stats, FLOORS, 0.5)] == ["paragraph spacing -> 0"]


def test_ladder_never_grows_margins_or_font():
    assert list(ladder(DocStats(spacing=False, line_spacing=1.0, margin_in=0.25, font_pt=9.5), FLOORS, 0.5)) == []


# ---------------------------------------------------------------- drops


def test_drop_bullet_prefers_last_project_and_never_empties():
    ctx = _ctx(project_bullets=(2, 1, 2))
    assert "P2" in drop_lowest_bullet(ctx)
    assert [len(p["bullets"]) for p in ctx["projects"]] == [2, 1, 1]
    assert "P0" in drop_lowest_bullet(ctx)
    assert drop_lowest_bullet(ctx) is None  # every project is down to one bullet


def test_drop_project_keeps_at_least_one():
    ctx = _ctx(project_bullets=(1, 1))
    assert "P1" in drop_lowest_project(ctx)
    assert drop_lowest_project(ctx) is None


# ---------------------------------------------------------------- fit loop


def _fake_measure(pages_for):
    calls = []

    def measure(ctx, adj):
        calls.append((copy.deepcopy(ctx), adj))
        return pages_for(ctx, adj), DocStats(spacing=True, line_spacing=1.15, margin_in=0.75, font_pt=11.0)

    return measure, calls


def test_fit_stops_as_soon_as_one_page(tmp_path):
    measure, calls = _fake_measure(lambda c, a: 1 if a.line_spacing else 2)
    res = fit_one_page(_ctx(), tmp_path / "x.docx", tmp_path / "x.pdf", measure=measure)
    assert res.fits and res.steps == ["paragraph spacing -> 0", "line spacing -> 1.0"] and res.dropped == []
    assert len(calls) == 3


def test_fit_drops_bullet_then_project(tmp_path):
    def pages(c, a):
        n_proj = len(c["projects"])
        n_bul = sum(len(p["bullets"]) for p in c["projects"])
        return 1 if (n_proj == 2 and n_bul <= 4) else 2

    measure, _ = _fake_measure(pages)
    res = fit_one_page(_ctx(), tmp_path / "x.docx", tmp_path / "x.pdf", measure=measure)
    assert res.fits and len(res.dropped) == 2
    assert res.dropped[0].startswith("bullet") and res.dropped[1].startswith("project")


def test_fit_gives_up_and_reports(tmp_path):
    measure, _ = _fake_measure(lambda c, a: 2)
    res = fit_one_page(_ctx(), tmp_path / "x.docx", tmp_path / "x.pdf", measure=measure, max_drops=3)
    assert not res.fits and res.pages == 2 and len(res.dropped) == 3


# ---------------------------------------------------------------- templates on both sample layouts


def test_ds_template_renders_master_text_exactly(ds, tmp_path):
    out = tmp_path / "r.docx"
    render.render_docx(master_context(ds.master), out)
    assert _texts(out) == _texts(ds.master)
    xml = Document(str(out)).element.xml
    assert "{%" not in xml and "{{" not in xml


def test_analytics_template_without_projects(analytics, tmp_path):
    ctx = master_context(analytics.master)
    assert ctx["projects"] == [] and ctx["skill_keys"] == ["row1"]
    out = tmp_path / "r.docx"
    render.render_docx(ctx, out)
    assert _texts(out) == _texts(analytics.master)


def test_detected_layout_of_ds_resume(ds):
    ctx = master_context(ds.master)
    assert ctx["skill_keys"] == ["languages", "machine_learning", "data_tools"]
    assert [p["title"] for p in ctx["projects"]] == ["Customer Churn Predictor", "Retail Demand Forecasting"]
    assert [p["term"] for p in ctx["projects"]] == ["jordan-lee/churn", "jordan-lee/forecast"]
    assert [p["term_link"] for p in ctx["projects"]] == ["https://github.com/jordan-lee/churn", "https://github.com/jordan-lee/forecast"]
    assert [len(p["bullets"]) for p in ctx["projects"]] == [3, 3]
    assert [len(e["bullets"]) for e in ctx["experience"]] == [3, 2]


def test_apply_adjust_on_rendered_resume(ds, tmp_path):
    out = tmp_path / "r.docx"
    render.render_docx(master_context(ds.master), out)
    base = render.body_font_pt(Document(str(out)))
    render.apply_adjust(out, Adjust(zero_spacing=True, margin_in=0.2, font_pt=base - 0.5))
    doc = Document(str(out))
    assert not render.has_paragraph_spacing(doc)
    assert render.body_font_pt(doc) == base - 0.5
    assert render.max_margin_in(doc) <= 0.21


def test_template_escapes_special_characters(ds, tmp_path):
    ctx = master_context(ds.master)
    ctx["projects"][0]["bullets"][0]["segs"] = [{"t": 'C++ & <Rust> "quoted"', "b": True}]
    out = tmp_path / "r.docx"
    render.render_docx(ctx, out)
    assert 'C++ & <Rust> "quoted"' in "".join(_texts(out))


# ---------------------------------------------------------------- links


def test_configured_link_in_bullet(ds, tmp_path):
    from jobpipe.segments import linkify

    ctx = master_context(ds.master)
    plain = [{"t": "Explained drivers in a ", "b": False}, {"t": "Streamlit", "b": True}, {"t": " report", "b": False}]
    ctx["projects"][0]["bullets"][1]["segs"] = linkify(plain, {"Streamlit": "https://streamlit.io/"})
    out = tmp_path / "r.docx"
    render.render_docx(ctx, out)
    assert _links(out)["Streamlit"] == "https://streamlit.io/"


def test_linkify_splits_occurrences():
    from jobpipe.segments import linkify

    segs = linkify([{"t": "on pwn.college and pwn.college.", "b": False}], {"pwn.college": "https://pwn.college/"})
    assert [s["t"] for s in segs] == ["on ", "pwn.college", " and ", "pwn.college", "."]
    assert [bool(s.get("link")) for s in segs] == [False, True, False, True, False]
    assert linkify([{"t": "nothing here", "b": True}], {"x.y": "u"}) == [{"t": "nothing here", "b": True}]
