import csv
import json

import pytest
from openpyxl import load_workbook

from jobpipe import lists
from jobpipe.config import path

URL = "https://jobright.ai/jobs/info/68f0a1b2c3d4e5f6a7b8c9d0"


def _row(**kw):
    base = {"company": "Acme", "role": "Engineer", "location": "Remote",
            "apply_url": "https://jobs.example.com/apply?a=1&b=\"x\"", "jobright_url": URL,
            "folder": r"2026-10-15\Acme_Engineer_1", "resume": "sde", "must_have_coverage": "50%",
            "new_project": "", "flags": ""}
    return {**base, **kw}


def _csv(f):
    return list(csv.DictReader(open(f, encoding="utf-8-sig")))


@pytest.fixture(autouse=True)
def both_formats(monkeypatch, request):
    """Most tests here check the CSV too; the default (xlsx only) has its own test."""
    if "xlsx_only" not in request.keywords:
        monkeypatch.setattr(lists, "formats", lambda: ["xlsx", "csv"])


@pytest.mark.xlsx_only
def test_default_writes_only_xlsx(tmp_path):
    assert lists.formats() == ["xlsx"]
    lists.append(tmp_path, "today", _row())
    assert (tmp_path / "_today.xlsx").exists() and not (tmp_path / "_today.csv").exists()
    assert (tmp_path / ".today.jsonl").exists()


def test_rows_go_to_ledger_csv_and_xlsx(tmp_path):
    lists.append(tmp_path, "skipped", {"company": "A", "role": "R", "jobright_url": URL, "reason": "x"})
    lists.append(tmp_path, "skipped", {"company": "B", "role": "R", "jobright_url": URL, "reason": "y"})
    assert [json.loads(ln)["company"] for ln in (tmp_path / ".skipped.jsonl").read_text().splitlines()] == ["A", "B"]
    assert [r["company"] for r in _csv(tmp_path / "_skipped.csv")] == ["A", "B"]
    ws = load_workbook(tmp_path / "_skipped.xlsx").active
    assert [c.value for c in ws[1]] == lists.SKIPPED_FIELDS and ws.max_row == 3


def test_csv_links_and_formula_guard(tmp_path):
    lists.append(tmp_path, "today", _row(company="=cmd|' /C calc'!A0", location="-Remote"))
    r = _csv(tmp_path / "_today.csv")[0]
    assert r["company"].startswith("'=") and r["location"] == "'-Remote"  # never evaluated
    assert r["apply_url"] == '=HYPERLINK("https://jobs.example.com/apply?a=1&b=%22x%22","Original Job Posting")'
    assert r["jobright_url"] == f'=HYPERLINK("{URL}","Open Jobright Listing")'
    assert r["folder"] == f'=HYPERLINK("{path("applications") / _row()["folder"]}","Open application directory")'
    long_url = "https://x.example/" + "a" * 300
    assert lists.excel_link(long_url, "x") == long_url  # over Excel's limit: plain text


def test_xlsx_has_styled_links_and_safe_text(tmp_path):
    lists.append(tmp_path, "today", _row(company="=1+1"))
    ws = load_workbook(tmp_path / "_today.xlsx").active
    header = [c.value for c in ws[1]]
    cell = {k: ws.cell(row=2, column=header.index(k) + 1) for k in header}
    assert cell["apply_url"].value == "Original Job Posting" and cell["apply_url"].hyperlink.target.startswith("https://jobs.example.com")
    assert cell["jobright_url"].value == "Open Jobright Listing" and cell["jobright_url"].hyperlink.target == URL
    assert cell["folder"].value == "Open application directory"
    assert cell["folder"].hyperlink.target == str(path("applications") / _row()["folder"])
    assert cell["apply_url"].font.u == "single" and cell["apply_url"].font.color.rgb.endswith("0563C1")
    assert cell["company"].value == "=1+1" and cell["company"].data_type == "s"  # text, not a formula
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref


def test_imports_csv_written_by_older_version(tmp_path):
    old = tmp_path / "_today.csv"
    with open(old, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=lists.TODAY_FIELDS)
        w.writeheader()
        w.writerow(_row(apply_url='=HYPERLINK("https://a.example/x","Apply")', company="Old Co"))
        w.writerow(_row(company="Plain Co"))  # the very first format: plain URLs
    lists.append(tmp_path, "today", _row(company="New Co"))
    rows = lists.load_rows(tmp_path, "today")
    assert [r["company"] for r in rows] == ["Old Co", "Plain Co", "New Co"]
    assert rows[0]["apply_url"] == "https://a.example/x"


def test_locked_file_warns_and_catches_up(tmp_path, monkeypatch):
    warnings = []
    real = lists.write_xlsx

    def locked(*a):
        raise PermissionError("in use")

    monkeypatch.setattr(lists, "write_xlsx", locked)
    lists.append(tmp_path, "today", _row(company="A"), warn=warnings.append)
    lists.append(tmp_path, "today", _row(company="B"), warn=warnings.append)
    assert len(warnings) == 1 and "_today.xlsx is open" in warnings[0]  # warned once, run continues
    assert [r["company"] for r in _csv(tmp_path / "_today.csv")] == ["A", "B"]
    monkeypatch.setattr(lists, "write_xlsx", real)
    assert lists.refresh(tmp_path, warn=warnings.append)
    assert load_workbook(tmp_path / "_today.xlsx").active.max_row == 3
