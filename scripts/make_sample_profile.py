"""Build the fictional sample profile in examples/sample_profile/ (used by the tests and as a
worked example for setting up a real profile). Run: uv run python scripts/make_sample_profile.py

Two resumes with deliberately different layouts, to exercise the generic template builder:
  ds         Title-case bold headings, "List Bullet" bullets, labeled skills rows, projects with a repo link
  analytics  "Heading 1" headings, one unlabeled skills row, no projects section
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "examples" / "sample_profile"


def _doc() -> Document:
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.6)
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(0)
    return d


def _rich(p, text: str) -> None:
    """Add runs to p; **x** is bold."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part:
            p.add_run(part).bold = i % 2 == 1


def _heading(d, text: str, style: str | None = None) -> None:
    if style:
        d.add_paragraph(text, style=style)
    else:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.add_run(text).bold = True


def _dated(d, left: str, right: str, bold_left: bool = True) -> None:
    p = d.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(left).bold = bold_left
    p.add_run("\t")
    p.add_run(right).bold = True


def _project(d, title: str, url: str) -> None:
    """Project header: bold title, then the GitHub repo as a hyperlink at the right margin."""
    p = d.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(title).bold = True
    p.add_run("	")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), d.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = url.removeprefix("https://github.com/")
    run.append(text)
    link.append(run)
    p._p.append(link)


def _bullet(d, text: str) -> None:
    _rich(d.add_paragraph(style="List Bullet"), text)


def _header(d) -> None:
    p = d.add_paragraph()
    p.alignment = 1
    r = p.add_run("Jordan Lee")
    r.bold = True
    r.font.size = Pt(18)
    c = d.add_paragraph("Chicago, IL | jordan.lee@example.com | 555-010-0199")
    c.alignment = 1


EXPERIENCE = [
    ("Data Analyst, Acme Retail, Chicago, IL", "June 2022 - Present", [
        "Cut weekly reporting time by **70%** by automating **SQL** extracts and **Python** notebooks into scheduled **Airflow** jobs.",
        "Lifted promotion ROI by **9%** for **120** stores by building an uplift model in **scikit-learn** with the marketing team.",
        "Reduced stockouts by **15%** by forecasting demand with **XGBoost** and publishing results to **Tableau** dashboards.",
    ]),
    ("Research Assistant, Lakeview University, Chicago, IL", "Aug 2020 - May 2022", [
        "Cleaned and merged **2 million** survey records with **pandas** and **R** for a public-health study.",
        "Built reproducible analysis pipelines in **Python** and **Git**, used by **6** researchers.",
    ]),
]


def make_ds(path: Path) -> None:
    d = _doc()
    _header(d)
    _heading(d, "Education")
    _dated(d, "Lakeview University, M.S. Data Science", "May 2022")
    _dated(d, "Riverside College, B.S. Statistics", "May 2020")
    _heading(d, "Skills")
    for row in ("**Languages:** Python, SQL, R",
                "**Machine Learning:** scikit-learn, XGBoost, PyTorch, statistical modeling",
                "**Data & Tools:** pandas, Spark, Airflow, Tableau, Git"):
        _rich(d.add_paragraph(), row)
    _heading(d, "Experience")
    for left, right, bullets in EXPERIENCE:
        _dated(d, left, right)
        for b in bullets:
            _bullet(d, b)
    _heading(d, "Projects")
    _project(d, "Customer Churn Predictor", "https://github.com/jordan-lee/churn")
    _bullet(d, "Raised churn recall by **18%** on **40,000** customers by training gradient-boosted models in **XGBoost** with class weighting.")
    _bullet(d, "Explained churn drivers to stakeholders through **SHAP** values in a **Streamlit** report.")
    _bullet(d, "Prepared customer features with **pandas** in a repeatable **Python** pipeline that feeds the model.")
    _project(d, "Retail Demand Forecasting", "https://github.com/jordan-lee/forecast")
    _bullet(d, "Lowered forecast error by **12%** MAPE by comparing **Prophet** and **LightGBM** models on store-level sales.")
    _bullet(d, "Automated backtesting across **52** weeks using **Python** and **pandas**.")
    _bullet(d, "Compared models on store-level sales in a reproducible **Python** notebook.")
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


def make_analytics(path: Path) -> None:
    d = _doc()
    _header(d)
    _heading(d, "SUMMARY", style="Heading 1")
    d.add_paragraph("Analyst who turns messy operational data into dashboards and decisions.")
    _heading(d, "SKILLS", style="Heading 1")
    d.add_paragraph("SQL, Python, Tableau, Excel, Looker, A/B testing, pandas")
    _heading(d, "EXPERIENCE", style="Heading 1")
    for left, right, bullets in EXPERIENCE:
        _dated(d, left, right)
        for b in bullets:
            _bullet(d, b)
    _heading(d, "EDUCATION", style="Heading 1")
    _dated(d, "Lakeview University, M.S. Data Science", "May 2022")
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


if __name__ == "__main__":
    make_ds(OUT / "resumes" / "ds" / "master_resume.docx")
    make_analytics(OUT / "resumes" / "analytics" / "master_resume.docx")
    print(f"wrote resumes under {OUT / 'resumes'}")
